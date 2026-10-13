# -*- coding: utf-8 -*-
"""Excel 输出：总览 / 信号明细 / 引擎指标 / 体检报告 / 评分卡 五张表。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from ..schema import Profile

HEAD_FILL = PatternFill("solid", fgColor="1F6FEB")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
BULL_FILL = PatternFill("solid", fgColor="FDECEB")   # 涨=红（A股习惯）
BEAR_FILL = PatternFill("solid", fgColor="E7F6EE")   # 跌=绿
FAIL_FILL = PatternFill("solid", fgColor="FEF4E7")

CONF_LABEL = {
    "bull_confluence": "看涨共振",
    "bear_confluence": "看跌共振",
    "divergence": "多空分歧",
    "no_confluence": "无有效共振",
}


def _head(ws, cols: list[str], widths: list[int]):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def _rows_overview(wb: Workbook, ps: list[Profile]):
    ws = wb.active
    ws.title = "画像总览"
    _head(ws, ["代码", "名称", "行业", "诊断基准日", "收盘", "共振结论", "共振分",
               "置信度", "趋势方向", "结构状态", "量价状态", "时效状态",
               "Wyckoff阶段", "道氏阶段", "防未来函数体检", "信号总数", "主要风险"],
          [10, 12, 18, 13, 9, 12, 8, 9, 14, 18, 20, 16, 13, 12, 15, 9, 46])
    for p in sorted(ps, key=lambda x: (-x.confluence.score, -x.scorecard.total)):
        guard_ok = all(g.passed for g in p.guard.values())
        n_sig = sum(len(r.signals) for r in p.engines.values())
        wk = p.engines.get("wyckoff")
        dw = p.engines.get("dow")
        ws.append([
            p.code, p.name, p.industry, p.as_of, p.close,
            CONF_LABEL.get(p.confluence.conclusion, p.confluence.conclusion),
            p.confluence.score, p.scorecard.total,
            p.confluence.trend, p.confluence.structure,
            p.confluence.volume_price, p.confluence.timeliness,
            (wk.metrics.get("phase") if wk and wk.ok else "-"),
            (dw.metrics.get("phase") if dw and dw.ok else "-"),
            "通过" if guard_ok else "未通过", n_sig,
            "；".join(p.risks) if p.risks else "无",
        ])
        r = ws.max_row
        if p.confluence.conclusion == "bull_confluence":
            ws.cell(r, 6).fill = BULL_FILL
        elif p.confluence.conclusion == "bear_confluence":
            ws.cell(r, 6).fill = BEAR_FILL
        if not guard_ok:
            ws.cell(r, 15).fill = FAIL_FILL
    return ws


def _rows_signals(wb: Workbook, ps: list[Profile]):
    ws = wb.create_sheet("信号明细")
    _head(ws, ["代码", "名称", "引擎", "信号类型", "方向", "状态",
               "形成日", "确认日", "失效日", "确认需K线", "价格", "强度", "证据"],
          [10, 12, 15, 20, 8, 12, 12, 12, 12, 11, 10, 8, 60])
    for p in ps:
        for r in p.engines.values():
            for s in r.signals:
                ev = "；".join(f"{k}={v}" for k, v in list(s.evidence.items())[:6])
                ws.append([p.code, p.name, s.engine, s.kind, s.direction, s.state,
                           s.formed_at, s.confirmed_at or "", s.invalidated_at or "",
                           s.confirm_bars, s.price, s.strength, ev])
    return ws


def _rows_metrics(wb: Workbook, ps: list[Profile]):
    ws = wb.create_sheet("引擎指标")
    _head(ws, ["代码", "名称", "POC", "VAH", "VAL", "收盘vsPOC", "POC箱量占比",
               "结构趋势", "最后结构事件", "RSI14", "最新背离", "背离状态",
               "K线形态", "结构形态已确认", "结构形态候选",
               "Wyckoff阶段", "假突破已确认", "放量无果", "道氏阶段", "回撤%", "量能偏向"],
          [10, 12, 10, 10, 10, 12, 13, 11, 26, 9, 26, 12, 24, 15, 15,
           13, 13, 11, 12, 9, 10])
    for p in ps:
        vp = p.engines.get("volume_profile")
        sm = p.engines.get("smc")
        dv = p.engines.get("divergence")
        pt = p.engines.get("pattern")
        wk = p.engines.get("wyckoff")
        dw = p.engines.get("dow")
        m_vp = vp.metrics if vp else {}
        m_sm = sm.metrics if sm else {}
        m_dv = dv.metrics if dv else {}
        m_pt = pt.metrics if pt else {}
        m_wk = wk.metrics if wk else {}
        m_dw = dw.metrics if dw else {}
        ret = m_dw.get("retracement") or {}
        vc = m_dw.get("volume_confirm") or {}
        last_candle = ""
        if pt and pt.signals:
            cs = [s for s in pt.signals if s.engine == "pattern" and s.kind in
                  ("doji", "hammer", "bullish_engulfing", "bearish_engulfing")]
            last_candle = f"{cs[-1].kind}@{cs[-1].formed_at}" if cs else ""
        ws.append([
            p.code, p.name,
            m_vp.get("poc"), m_vp.get("vah"), m_vp.get("val"),
            m_vp.get("close_vs_poc"), m_vp.get("poc_bin_volume_ratio"),
            m_sm.get("structure_trend"), m_sm.get("last_structure"),
            m_dv.get("rsi_last"), m_dv.get("latest"), m_dv.get("latest_state"),
            last_candle, m_pt.get("structure_confirmed"), m_pt.get("structure_candidate"),
            m_wk.get("phase"), m_wk.get("confirmed_break_events"),
            m_wk.get("effort_no_result_count"),
            m_dw.get("phase"), ret.get("retracement_pct"), vc.get("volume_bias"),
        ])
    return ws


def _rows_guard(wb: Workbook, ps: list[Profile]):
    ws = wb.create_sheet("防未来函数体检")
    _head(ws, ["代码", "名称", "引擎", "是否通过", "信号数", "重算不一致",
               "未来函数违规", "确认时间违规", "样本"],
          [10, 12, 16, 10, 9, 12, 13, 13, 60])
    for p in ps:
        for g in p.guard.values():
            ws.append([p.code, p.name, g.engine, "通过" if g.passed else "未通过",
                       g.total_signals, g.mismatched, g.lookahead_violations,
                       g.confirm_violations, " | ".join(g.samples)])
            if not g.passed:
                ws.cell(ws.max_row, 4).fill = FAIL_FILL
    return ws


def _rows_score(wb: Workbook, ps: list[Profile]):
    ws = wb.create_sheet("六维评分卡")
    dims = ["算法完整性", "信号可靠性", "工程可集成性", "测试与验证", "A股适配性", "许可证与维护"]
    _head(ws, ["代码", "名称"] + dims + ["加权总分"],
          [10, 12] + [13] * len(dims) + [11])
    for p in sorted(ps, key=lambda x: -x.scorecard.total):
        ws.append([p.code, p.name] +
                  [p.scorecard.dims.get(d) for d in dims] +
                  [p.scorecard.total])
    return ws


def write(ps: list[Profile], out: Path) -> Path:
    wb = Workbook()
    _rows_overview(wb, ps)
    _rows_signals(wb, ps)
    _rows_metrics(wb, ps)
    _rows_guard(wb, ps)
    _rows_score(wb, ps)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
